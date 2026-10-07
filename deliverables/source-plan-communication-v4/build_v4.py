from pathlib import Path
import json, math, re, unicodedata
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from lxml import etree

ROOT=Path(__file__).resolve().parent;OUT=ROOT.parent/'output';AS=ROOT/'assets'
P=Presentation(ROOT/'base_v3.pptx')
ORIGINAL_IDS=list(P.slides._sldIdLst)
N='002C49';O='FF6633';W='FFFFFF';PAPER='F7F7F3';PALE='EAF0F3';G='526777';L='D9E2E7';M='9DB0BD';BLUE='0C405C';FONT='Montserrat'
REGISTRY=[]

def box(s,x,y,w,h,fill,line=None,shape=MSO_SHAPE.RECTANGLE):
 a=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:a.line.color.rgb=RGBColor.from_string(line);a.line.width=Pt(1)
 else:a.line.fill.background()
 style=a._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
 if style is not None:a._element.remove(style)
 return a

def text(s,x,y,w,h,value,size=28,c=N,bold=False,align=None,leading=None):
 a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=a.text_frame;f.word_wrap=True;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
 for i,line in enumerate(value.split('\n')):
  q=f.paragraphs[0] if i==0 else f.add_paragraph();q.text=line;q.font.name="Nohemi" if size>=35 else FONT;q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(c)
  q.space_after=Pt(0);q.space_before=Pt(0);q.line_spacing=Pt(leading or size*1.23)
  if align is not None:q.alignment=align
 REGISTRY.append({'slide':len(P.slides),'text':value,'x':x*72,'y':y*72,'w':w*72,'h':h*72,'size':size})
 return a

def rule(s,x,y,w,c=L):box(s,x,y,w,.012,c)
def dot(s,x,y,r=.16,c=O):return box(s,x-r/2,y-r/2,r,r,c,shape=MSO_SHAPE.OVAL)
def line(s,x1,y1,x2,y2,c=L,width=1.2):
 a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));a.line.color.rgb=RGBColor.from_string(c);a.line.width=Pt(width)
 style=a._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
 if style is not None:a._element.remove(style)
 return a

def tag(s,x,y,label,fill=PALE,c=N,w=None):
 w=w or max(1.1,len(label)*.105+.35);box(s,x,y,w,.43,fill);text(s,x+.13,y+.09,w-.26,.26,label,14,c,True)

def footer(s,idx,dark=False,annex=False):
 cc=M if dark else G;rule(s,.9,10.65,18.2,BLUE if dark else L)
 text(s,.9,10.85,16,.23,'LEYTON FRANCE  /  PLAN DE COMMUNICATION  /  FY2026/27',12,cc)
 text(s,18.1,10.8,1,.3,f'{idx:02d}',16,O,True,PP_ALIGN.RIGHT)

def new(section,title='',subtitle=None,dark=False,notes=None):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(N if dark else W)
 s.shapes.add_picture(str(AS/f'logo{2 if dark else 3}.png'),Inches(.9),Inches(.55),width=Inches(1.85))
 text(s,4.0,.65,15,.32,section.upper(),15,M if dark else G,True,PP_ALIGN.RIGHT)
 if title:text(s,.9,1.5,18.2,1.56,title,46,W if dark else N,False,leading=53)
 if subtitle:text(s,.95,2.999,18.05,.95,subtitle,24,M if dark else G,leading=29)
 footer(s,len(P.slides),dark)
 if notes:s.notes_slide.notes_text_frame.text=notes
 return s

def takeaway(s,value,dark=False,y=9.52):
 rule(s,.9,y,18.2,BLUE if dark else L);text(s,.95,y+.22,18.03,.84,value,24,W if dark else N,True,leading=28)

def numbered(s,x,y,num,heading,body,w=8):
 text(s,x,y,.9,.65,num,38,O,True);text(s,x+1.1,y+.05,w-1.1,.92,heading,29,N,True)
 text(s,x+1.1,y+1.1,w-1.1,1.5,body,26,G)

# V4: reuse V3 drawing helpers and original slides as a read-only base.
import hashlib
from pptx.enum.dml import MSO_LINE_DASH_STYLE
BASE_HASH='358360fdd062e4a16ae34ac9aa7c6c90dbb3022b441d1cb91854c95b327578c6'
assert hashlib.sha256((ROOT/'base_v3.pptx').read_bytes()).hexdigest()==BASE_HASH
MAIN=[];EXTRA=[]
def page(section,title,subtitle=None,dark=False,notes=None,appendix=False):
 s=new(section,title,subtitle,dark,notes)
 (EXTRA if appendix else MAIN).append(P.slides._sldIdLst[-1])
 return s

def para_replace(s,old,value):
 for sh in s.shapes:
  if sh.has_text_frame and sh.text==old:
   p=sh.text_frame.paragraphs[0]
   if p.runs:
    p.runs[0].text=value
    for r in list(p.runs)[1:]:r._r.getparent().remove(r._r)
   else:p.text=value
   for extra in list(sh.text_frame.paragraphs)[1:]:extra._p.getparent().remove(extra._p)
   return
 raise ValueError(old)

def label(s,x,y,value,c=O,w=8):text(s,x,y,w,.42,value,18,c,True,leading=23)
def card(s,x,y,w,h,heading,body,dark=False,num=None):
 box(s,x,y,w,h,BLUE if dark else PAPER)
 if num:label(s,x+.3,y+.27,num,O,w-.6)
 text(s,x+.3,y+(.85 if num else .3),w-.6,.92,heading,30,W if dark else N,True,leading=35)
 text(s,x+.3,y+(1.95 if num else 1.48),w-.6,h-(2.13 if num else 1.66),body,26,M if dark else G,leading=33)

# 1 / cover
s=page('France • juillet 2026 – juin 2027','',dark=True,notes='V4 exécutive. Base V3 conservée, portefeuille inchangé. Les détails et sources sont en annexes. Priorisation et séquencement proposés : à arbitrer avec la capacité réelle.')
text(s,.9,2.05,12.6,3.65,'Plan de\nCommunication\nFrance',67,W,leading=77)
text(s,.92,6.45,13.6,1.5,'FY2026/27.',89,O,leading=100)
line(s,14.0,2.15,14.0,9.45,BLUE,1.5)
s.shapes.add_picture(str(AS/'logo1.png'),Inches(15.1),Inches(2.45),width=Inches(3.7))
text(s,15.0,6.7,4.05,2.45,'Pilotage.\nOrchestration.\nInfluence.',29,W,leading=40)
text(s,.97,9.16,12.7,.55,'Les choix de Communication France, au service de Leyton 2030.',28,W)

# 2 / diagnostic
s=page('01 / Diagnostic','Trois constats. Trois réponses pour FY2026/27.',subtitle='Un portefeuille pertinent ; une manière de le piloter à faire évoluer.')
label(s,1.0,3.9,'CONSTATS',G,8);label(s,10.5,3.9,'RÉPONSES',O,8)
rows=[('Plusieurs transformations\nstructurantes à accompagner.','Relier les actions à Leyton 2030\net séquencer les priorités.'),('Une logique encore centrée\nsur la production et la diffusion.','Piloter les effets recherchés ;\norchestrer messages et relais.'),('Un portefeuille important,\ndes canaux à mieux articuler.','Prioriser, mutualiser et capitaliser\npour éviter les doublons.')]
for i,(a,b) in enumerate(rows):
 y=4.6+i*1.55; text(s,1,y,.8,.65,f'0{i+1}',38,O,True)
 text(s,2.15,y+.05,7.7,1.25,a,29,N,True,leading=35)
 text(s,10.5,y+.05,8.2,1.08,b,28,G,leading=35)
 if i<2:rule(s,2.15,y+1.25,16.5)
takeaway(s,'Faire mieux, pas plus : moins de dispositifs isolés, davantage de cohérence et de capitalisation.')

# 3 / ambition
s=page('02 / Ambition','Piloter. Orchestrer. Influencer.',subtitle='Passer d’une logique de production et de diffusion à une logique de pilotage, d’orchestration et d’influence.')
for i,(h,b) in enumerate([('PILOTER','Choisir et séquencer\nles priorités.'),('ORCHESTRER','Faire travailler ensemble\nmessages, projets et canaux.'),('INFLUENCER','Contribuer à l’adoption,\nà la crédibilité et à la mobilisation.')]):
 x=.95+i*6.1;text(s,x,4.5,5.75,.65,h,37,O,True);rule(s,x,5.5,5.45)
 text(s,x,5.85,5.45,1.9,b,29,N,leading=36)
text(s,.95,8.6,18.1,.6,'PRIORISER  /  MUTUALISER  /  RÉUTILISER  /  FUSIONNER',25,N,True)
takeaway(s,'La Communication contribue aux résultats avec les équipes responsables ; elle ne les porte pas seule.')
s.notes_slide.notes_text_frame.text='Faire mieux, pas plus : ne pas tout lancer simultanément ; une production peut servir plusieurs canaux ; réutiliser interviews, vidéos et preuves ; fusionner les dispositifs redondants. Arrêter de produire pour remplir un canal. L’influence est une contribution, sans attribution exclusive des ventes, de l’adoption ou de la performance business à Communication.'

# 4 / narratives
s=page('03 / Leyton 2030','Quatre grands récits à rendre visibles.',subtitle='Les projets sont les preuves et dispositifs qui servent ces récits.')
items=[('Discipline financière & performance','PERFORMANCE','Faire mieux, mieux choisir,\nrendre la performance compréhensible.'),('Montée en gamme & leadership marché','MONTÉE EN GAMME','Renforcer la perception d’expertise,\nde valeur et de crédibilité.'),('Digital, IA & outils','DIGITAL / IA / OUTILS','Accompagner la transformation\ndes métiers et des usages.'),('Culture, équipes & collectif','CULTURE & COLLECTIF','Renforcer compréhension, engagement,\ncollaboration et culture commune.')]
for i,(p,h,b) in enumerate(items):
 x=.95+(i%2)*9.3;y=3.85+(i//2)*2.75
 box(s,x,y,8.9,2.45,PAPER);label(s,x+.3,y+.25,p,G,8.3)
 text(s,x+.3,y+.78,8.3,.7,h,32,N,True)
 text(s,x+.3,y+1.5,8.3,.85,b,25,G,leading=31)
takeaway(s,'Une action peut servir plusieurs récits. Aucun projet n’est créé pour remplir une case.')

# 5 / matrix
s=page('04 / Contributions','Trois piliers au service de Leyton 2030.',subtitle='Contributions principales — les piliers se croisent, ils ne se succèdent pas.')
xs=[.95,7.1,11.1,15.1];widths=[5.8,3.65,3.65,3.85]
for x,w,h in zip(xs,widths,['LEYTON 2030','CHANGE','BUSINESS','PEOPLE']):text(s,x,3.85,w,.48,h,25,N,True)
rule(s,.95,4.45,18.1,N)
mat=[('Discipline financière\n& performance','Lisibilité de\nla performance','Preuves de\nvaleur','Repères\ncollectifs'),('Montée en gamme\n& leadership marché','Offres et expertises\nmobilisables','Expertise\net crédibilité','Collaborations\net expertises'),('Digital, IA & outils','Adoption\ndes usages','Valeur pour\nle client','Accès aux\nressources'),('Culture, équipes\n& collectif','Sens des\ntransformations','—','Engagement\net culture')]
for j,row in enumerate(mat):
 y=4.7+j*1.06
 for k,v in enumerate(row):text(s,xs[k],y,widths[k],.9,v,25,N if k==0 else G,k==0,leading=29)
 if j<3:rule(s,.95,y+.95,18.1)
text(s,.95,9.22,18.1,.62,'CHANGE : comprendre / adopter  •  BUSINESS : preuve / crédibilité  •  PEOPLE : information / engagement / collectif',20,N,True,leading=25)
s.notes_slide.notes_text_frame.text='La matrice décrit les contributions principales, pas une sélection de projets ni une roadmap. Les projets concrets figurent ensuite. Culture × Business : aucun projet autonome ; une action peut contribuer à plusieurs priorités. Les trois piliers conservent les définitions validées : Change faire comprendre et adopter ; Business renforcer preuve, expertise et crédibilité ; People renforcer information, engagement, culture et collectif.'

# 6 / hierarchy
s=page('05 / Choix FY2026/27','Six ensembles structurants. Un socle à maintenir.',subtitle='Hiérarchisation proposée ; le séquencement reste à arbitrer dans la capacité réelle.')
label(s,.95,3.9,'PRIORITÉS STRUCTURANTES',O,12)
prior=['Adoption de l’IA','Toolbox / adoption des outils','Business Partner & cross-sell','Preuves & montée en gamme','Performance','Culture, engagement & image corporate']
for i,v in enumerate(prior):
 x=.95+(i%2)*6.35;y=4.55+(i//2)*1.26
 box(s,x,y,6.02,1.16,PAPER);text(s,x+.22,y+.19,.55,.6,f'{i+1}',29,O,True)
 text(s,x+.9,y+.12,4.86,.94,v,25,N,True,leading=29)
box(s,13.85,3.9,5.25,2.43,PALE);label(s,14.15,4.18,'SOCLE / RUN',N,4.65)
text(s,14.15,4.85,4.65,1.42,'Rituels, rendez-vous,\nRSE et canaux permanents.',26,G,leading=32)
box(s,13.85,6.58,5.25,2.63,PAPER);label(s,14.15,6.87,'CONDITIONNEL / TBD',O,4.65)
text(s,14.15,7.55,4.65,1.53,'Selon les lancements\net les opportunités.\nCadrages à confirmer.',25,G,leading=30)
text(s,.95,8.7,12.25,.75,'Écosystème interne à transformer : Neoscreen / intranet.',25,N,True)
takeaway(s,'Ne pas tout lancer simultanément. Les six ensembles ne représentent pas six projets de poids égal.')
s.notes_slide.notes_text_frame.text='Six regroupements de lecture, sans renommer les projets. Les priorités sont une proposition stratégique ; aucune capacité chiffrée ni ordre de lancement arbitraire. Culture/engagement/image combine dispositifs structurants et socle, détaillés slide 9. Écosystème interne traité séparément : Neoscreen déploiement dès octobre puis run ; intranet revue de production, possible changement de prestataire, planning après arbitrage. Conditionnel : nouvelles offres selon lancements, captations selon opportunités, vidéo Groupe, certains événements, cadrage Workelo et Business Partner.'

# 7 / Change
s=page('06 / Change','Faire comprendre. Équiper. Accompagner les usages.',subtitle='Des dispositifs complémentaires, pilotés avec les équipes responsables du fond.')
rows=[('Adoption de l’IA','Un parcours collaborateur,\npas un seul temps fort.','Tous collaborateurs'),('Toolbox / adoption des outils','Un hub et des relais multicanaux ;\nun standard de lancement.','Collaborateurs utilisateurs'),('Business Partner & cross-sell','Équiper par pôle ; rendre visibles\net capitaliser les collaborations.','Commerciaux / consultants'),('Performance','Offres qui performent et relais\ndes chiffres validés post-All Staff.','Équipes / managers')]
for i,(h,b,t) in enumerate(rows):
 y=3.95+i*1.24;text(s,.95,y,5.85,.91,h,28,N,True,leading=33)
 text(s,7.1,y,7.7,.98,b,25,G,leading=30);text(s,15.2,y+.04,3.75,.84,t,22,O,True,leading=27)
 if i<3:rule(s,.95,y+1.08,18.1)
text(s,.95,9.03,18.1,.85,'Neoscreen : déploiement à partir d’octobre.  Intranet : planning après revue de production / arbitrage prestataire.',24,N,True,leading=30)
s.notes_slide.notes_text_frame.text='Kit Business Partner — nom de travail : périmètre et owner du fond à valider, Communication porte expérience/branding/cohérence et Marketing/métiers le fond commercial. Cross-sell = sourcing → sélection → interview → validation → diffusion → capitalisation ; minimum 4 vidéos/exercice. Zoom IA à confirmer et ne bloque pas le programme. Hub IA sur intranet après lancement ; relais disponibles avant. Performance distincte des Top Performers Business. Aucun lancement intranet en octobre.'

# 8 / Business
s=page('07 / Business','Produire des preuves de valeur.',subtitle='Renforcer l’expertise, la crédibilité et la montée en gamme — avec Marketing, Business et Customer Success.')
box(s,.95,3.9,11.0,4.92,PAPER);label(s,1.3,4.25,'PREUVES & MONTÉE EN GAMME',O,10.3)
text(s,1.3,4.98,10.25,1.1,'Problème client → accompagnement Leyton\n→ résultat / valeur créée.',30,N,True,leading=37)
text(s,1.3,6.4,10.25,1.5,'Témoignages premium • outils orientés client\n• témoignages Tech + expérience client.\nCibles : clients / prospects, avec relais internes.',26,G,leading=33)
text(s,12.6,4.12,6.45,.69,'TOP PERFORMERS',30,N,True)
text(s,12.6,5.07,6.45,1.50,'Performance individuelle / collective.\nOct. → mi-déc. ; fév. → avr.',26,G,leading=33)
rule(s,12.6,6.72,6.45)
label(s,12.6,7.0,'ACTIVATION CONDITIONNELLE',O,6.45)
text(s,12.6,7.6,6.45,1.3,'Nouvelles offres selon lancements.\nCaptations selon opportunités.\nVidéo corporate : calendrier Groupe.',25,G,leading=31)
takeaway(s,'Marketing / Business / CS : fond commercial. Communication : récit, formats, cohérence et amplification.')
s.notes_slide.notes_text_frame.text='Témoignages principalement vidéo PME/ETI/Grands Comptes lorsque possible. Salons = opportunité de captation, pas une campagne autonome. Oryx modèle de valorisation d’outil ; Leyton For Me si cas démontrable. Nouvelles offres : expert/vidéo/fiche/preuve/interne/externe avec Marketing selon lancement réel. La vidéo corporate Groupe montre expertise, montée en gamme, preuves clients, Tech, IA et outils propriétaires ; calendrier Groupe à confirmer. Les offres qui performent restent Change, distinctes des Top Performers.'

# 9 / People
s=page('08 / People & image corporate','Mobiliser le collectif. Donner à voir la culture Leyton.',subtitle='Une priorité de culture et d’image, appuyée sur des dispositifs RH et un socle régulier.')
for i,(h,b) in enumerate([('CULTURE & ENGAGEMENT','CARE : résultats → enseignements → actions RH.\nCulture Leyton : Kudos, initiatives, collaborations.\nWorkelo : refonte J−60 → J+30, cadrage à confirmer.'),('IMAGE CORPORATE','Leadership au féminin : 4 capsules nov. → fév.\nPuis un mash-up autour du 8 mars.\nLinkedIn : cinq territoires, selon calendrier joint.\nRSE : interne, puis sélection à l’externe.')]):
 x=.95+i*9.3;box(s,x,3.9,8.9,3.6,PAPER);label(s,x+.3,4.2,h,O,8.3)
 text(s,x+.3,4.93,8.3,2.45,b,25,G,leading=34)
label(s,.95,7.8,'SOCLE : RITUELS & RENDEZ-VOUS',N,12)
text(s,.95,8.36,18.1,.9,'All Staff • La Leytonienne • rendez-vous collaborateurs • RSE • temps forts RH / collectifs.\nCibles : collaborateurs et managers ; image corporate à l’externe.',25,G,leading=32)
takeaway(s,'Le sport est un territoire de marque : collaborateurs / CSE et Sport Business. Un événement = maximum un post.')
s.notes_slide.notes_text_frame.text='4 All Staff octobre/janvier/avril/fin exercice ; Leytonienne S1 fin janvier, S2 préparée mai–juin publiée fin juillet hors FY. Minimum un RDV collaborateurs et une activation RSE interne/mois. RSE externe seulement sélection. Christmas Party et séminaire fin exercice, date à confirmer. CARE déjà lancé, ne pas annoncer nouveau lancement ; Communication relie résultats et actions RH, mobilité/formation/développement/parcours. Cascade managériale grands sujets uniquement. Sport CSE/collaborateurs HYROX/semi-marathon/triathlon ; Sport Business hospitalités et événements sélectionnés, aucune couverture systématique. Le calendrier LinkedIn est la source de vérité ; ni calendrier complet du run RH/RSE interne ni plan Marketing.'

# 10 / channels
s=page('09 / Écosystème','Un rôle par canal. Une production, plusieurs usages.',subtitle='Orchestrer les contenus selon l’effet recherché ; ne pas produire pour remplir un canal.')
label(s,.95,3.95,'INTERNE',O,8.5);label(s,10.3,3.95,'EXTERNE CORPORATE FRANCE',O,8.5)
left=[('Intranet','Hub / ressources / capitalisation — après lancement.'),('Neoscreen','Amplification en agences — après déploiement.'),('Teams / email','Sourcing et activation ciblée selon le besoin.'),('All Staff / Leytonienne','Alignement collectif / éditorialisation.')]
right=[('LinkedIn','Canal corporate prioritaire ; calendrier éditorial.'),('Instagram','Activation selon contenus visuels / événements.'),('YouTube','Hébergement et valorisation des vidéos structurantes.')]
for x,items in [(.95,left),(10.3,right)]:
 for i,(h,b) in enumerate(items):
  y=4.62+i*1.17;text(s,x,y,8.7,.45,h,27,N,True);text(s,x,y+.5,8.7,.62,b,23,G,leading=28)
text(s,10.3,8.55,8.7,.62,'LinkedIn business : Marketing.',25,N,True)
takeaway(s,'Projet de déploiement / lancement ≠ alimentation continue. Les trois piliers nourrissent les canaux disponibles.')

# timeline helpers (native PowerPoint)
MONTHS=['JUIL.','AOÛT','SEPT.','OCT.','NOV.','DÉC.','JAN.','FÉV.','MARS','AVR.','MAI','JUIN']
TX=5.85;TW=13.2;CW=TW/12

def timehead(s,y=3.85):
 for k,m in enumerate(MONTHS):
  text(s,TX+k*CW,y,CW,.4,m,17,N,True,PP_ALIGN.CENTER)
  if k==1:box(s,TX+k*CW,y+.56,CW,4.5,PAPER)

def bar(s,y,a,b,c=N,dashed=False):
 if dashed:
  v=line(s,TX+a*CW,y,TX+b*CW,y,c,2.0);v.line.dash_style=MSO_LINE_DASH_STYLE.DASH
 else:box(s,TX+a*CW,y-.06,(b-a)*CW,.12,c)

def mark(s,y,m,c=O):dot(s,TX+(m+.5)*CW,y,.2,c)

# 11 / executive roadmap
s=page('10 / Temporalité','Une roadmap exécutive. Les fenêtres qui comptent.',subtitle='Juillet 2026 → juin 2027 • août et vacances : rythme réduit • aucun score de charge.')
timehead(s)
for y in [4.67,6.19,7.71]:rule(s,.95,y,18.1)
label(s,.95,4.87,'CHANGE',O,4.7);text(s,.95,5.43,4.65,.63,'IA / outils / BP & cross-sell',22,N,True)
bar(s,5.02,0,12,N,dashed=True)
text(s,TX+.12,5.14,12.9,.52,'Programmes : séquencement après arbitrage priorités / capacité',19,G)
mark(s,5.87,3);text(s,TX+3.85*CW,5.59,8.7,.54,'Neoscreen : à partir d’octobre',20,N,True)
label(s,.95,6.39,'BUSINESS',O,4.7);text(s,.95,6.97,4.65,.55,'Fenêtres de performance',22,N,True)
bar(s,6.73,3,5.5,O);bar(s,6.73,7,10,O)
text(s,TX+3.05*CW,6.94,3.0,.6,'Oct. → mi-déc.',20,N,True)
text(s,TX+7.1*CW,6.94,3.2,.6,'Fév. → avr.',20,N,True)
label(s,.95,7.91,'PEOPLE',O,4.7);text(s,.95,8.5,4.65,.67,'Rituels / série Leadership',22,N,True)
for k in [3,6,9,11]:mark(s,8.08,k,N)
text(s,TX+.12,8.28,12.9,.48,'All Staff : oct. / jan. / avr. / fin d’exercice*',19,G)
for k in [4,5,6,7]:mark(s,8.89,k,O)
mark(s,8.89,8,O)
text(s,TX+.12,9.08,12.9,.43,'Leadership : capsules nov.–fév. ; mash-up autour du 8 mars',19,G)
text(s,.95,9.65,18.1,.85,'Fils rouges : RDV / RSE internes mensuels ; LinkedIn corporate. Leytonienne : S1 fin janvier ; S2 préparée mai–juin, publiée fin juillet hors FY.\nPoint = temps fort ; barre = fenêtre ; pointillé = séquencement à confirmer. *All Staff fin d’exercice : date à confirmer.',18,G,leading=23)
s.notes_slide.notes_text_frame.text='Roadmap exécutive, pas calendrier de production. Point = temps fort ; barre = fenêtre validée ; pointillé = séquencement à définir et non engagement sur toute la période. Les programmes IA/Toolbox/Kit/cross-sell n’ont pas de dates de lancement validées. Repères : les offres qui performent et Top Performers octobre à mi-décembre et février à avril ; Neoscreen campagne à partir d’octobre. All Staff fin exercice positionné dans la fenêtre de juin, date exacte à confirmer. Fils rouges RDV et RSE mensuels et LinkedIn corporate toute année ; détaillés en annexe. Leytonienne S1 fin janvier ; S2 préparation mai–juin, publication fin juillet hors FY. Zoom IA distinct du post LinkedIn IA 22 octobre.'

# 12 / light governance & measurement
s=page('11 / Pilotage','Des responsabilités claires. Des effets suivis.',subtitle='Le pilotage relie les priorités, les validations et les données disponibles.')
label(s,.95,3.95,'GOUVERNANCE LÉGÈRE',O,8.7);label(s,10.3,3.95,'MESURE PAR GRAND PROGRAMME',O,8.7)
text(s,.95,4.7,8.55,2.05,'Owner du fond / contributeurs\nRôle Communication\nArbitrage nécessaire',30,N,True,leading=47)
text(s,.95,7.17,8.55,1.62,'Direction : priorités et arbitrages.\nÉquipes expertes : fond et validation.\nCommunication : orchestration / capitalisation.',25,G,leading=34)
text(s,10.3,4.7,8.55,2.05,'1 effet recherché\n1 ou 2 indicateurs\nUne source de données',30,N,True,leading=47)
text(s,10.3,7.17,8.55,1.6,'Adoption, participation, usage, preuves.\nSources à mobiliser avec les équipes.\nAucune cible chiffrée ajoutée.',25,G,leading=34)
takeaway(s,'Les volumes de production sont des repères d’activité, pas une preuve d’impact. Détails par programme en annexes.')
s.notes_slide.notes_text_frame.text='Owners opérationnels, validations et rythme de pilotage à confirmer. IA : Direction sens, DSI outils/anonymisation, Juridique cadre, RH Formation compétences, métiers cas. Marketing/Business/CS propriétaires du fond commercial. Groupe propriétaire des projets et calendriers Groupe. Périmètre Sport Business distinct. Indicateurs proposés à valider selon accès aux données ; aucun KPI réalisé ni objectif numérique inventé. Les repères fournis sont minimum 4 vidéos cross-sell/exercice, 1 RDV collaborateurs/mois, 1 activation RSE/mois, 4 All Staff ; ils ne prouvent pas l’impact.'

# 13 / decisions
s=page('12 / Décisions attendues','Les arbitrages qui rendent le plan exécutable.',dark=True,subtitle='Valider les choix, leur séquencement et les conditions de mise en œuvre.')
cols=[('PRIORITÉS & CAPACITÉ','Quels ensembles lancer en premier ?\nQue peut absorber l’équipe ?\nQue séquencer ou différer ?'),('DÉPENDANCES & CADRAGES','Intranet : revue / prestataire.\nIA : date du Zoom,\nsans bloquer le parcours.\nBusiness Partner : périmètre / owner.\nWorkelo : cadrage / temporalité.'),('CALENDRIERS & PILOTAGE','Vidéo corporate : calendrier Groupe.\nSéminaire : date à confirmer.\nOwners, validations et rythme\nde pilotage à arrêter.')]
for i,(h,b) in enumerate(cols):
 x=.95+i*6.1;box(s,x,3.95,5.75,4.96,BLUE);label(s,x+.28,4.29,h,O,5.19)
 text(s,x+.28,5.14,5.25,3.75,b,26,W,leading=37)
takeaway(s,'L’enjeu : arrêter un plan priorisé et pilotable dans la capacité réelle de Communication France.',True)

# Preserve selected original slides, with limited required content updates.
KEEP=[6,21,7,22,9,23,12,24,15,25,26]
ANNEX_IDS=[ORIGINAL_IDS[n-1] for n in KEEP]
retained=[P.slides[n-1] for n in KEEP]
for a,(n,s) in enumerate(zip(KEEP,retained),1):
 for sh in s.shapes:
  if sh.has_text_frame and sh.top<Inches(1) and sh.left>Inches(3):
   para_replace(s,sh.text,f'ANNEXE A{a:02d} / DÉTAILS DU PORTEFEUILLE');break
 if n==6:
  for sh in s.shapes:
   if sh.has_text_frame and sh.text in ['Comment passer\nà l’action ?', 'Formations, hub\nChallenge, pack métiers']:
    sh.width=Inches(3.35);sh.height=Inches(1.55 if sh.text.startswith('Comment') else 1.35)
 if n==23:
  para_replace(s,'Expertise, montée en gamme, preuves clients, Tech, IA,\noutils propriétaires ; calendrier Groupe TBD.','Expertise ; montée en gamme ; preuves clients ; Tech / IA.\nOutils propriétaires ; calendrier Groupe TBD.')
  for sh in s.shapes:
   if sh.has_text_frame and sh.text.startswith('Expertise ; montée en gamme'):
    sh.height=Inches(1.23);sh.top=Inches(8.35)
 if n==21:
  para_replace(s,'Hub intranet après lancement ; relais disponibles avant. Zoom IA ne remplace pas le parcours.','Hub après lancement ; relais disponibles avant. Zoom IA à confirmer, sans bloquer le parcours.')
 if n==22:
  label(s,.95,3.28,'KIT BUSINESS PARTNER — NOM DE TRAVAIL',G,17.9)
 if n==24:
  para_replace(s,'Parcours, management, conseils, vision.\nCapsules régulières + mash-up externe le 8 mars.','4 capsules : novembre, décembre, janvier, février.\nMash-up des quatre autour du 8 mars ; angles en calendrier.')
 if n==15:
  para_replace(s,'Revue de production / prestataire ; planning TBD.','Planning après revue de production / arbitrage prestataire.')
 if n==15:
  for sh in s.shapes:
   if sh.has_text_frame and sh.text=='Planning après revue de production / arbitrage prestataire.':sh.height=Inches(.98);sh.top=Inches(3.8)
 s.notes_slide.notes_text_frame.text+='\nV4 : détail conservé en annexe. Les calendriers non validés restent à confirmer. Le calendrier LinkedIn actualisé prévaut sur toute formulation générique antérieure.'

# A12 / detailed projects roadmap
s=page('Annexe A12 / Roadmap détaillée','Projets : fenêtres validées et dépendances.',subtitle='Pointillés = temporalité à arbitrer, pas une période de production engagée.',appendix=True)
timehead(s)
rr=[('CHANGE / IA & Toolbox','seq'),('CHANGE / BP & cross-sell','seq'),('CHANGE / Neoscreen','neo'),('Les offres qui performent','perf'),('BUSINESS / Preuves clients','seq'),('BUSINESS / Top Performers','perf')]
for i,(h,k) in enumerate(rr):
 y=4.65+i*.74;text(s,.95,y-.2,4.6,.62,h,22,N,True,leading=26)
 if k=='seq':bar(s,y,0,12,G,True)
 elif k=='neo':mark(s,y,3);bar(s,y,3.4,12,N,True)
 else:bar(s,y,3,5.5,O);bar(s,y,7,10,O)
 rule(s,.95,y+.33,18.1)
text(s,.95,9.21,18.1,.85,'Intranet : après revue / arbitrage prestataire.  Nouvelles offres : selon lancements.\nVidéo corporate : calendrier Groupe.  Zoom IA : date à confirmer, parcours non bloqué.',22,G,leading=29)
s.notes_slide.notes_text_frame.text='Reprise et simplification des données de la roadmap V3, sans nouvelles dates. IA, Toolbox, Kit et cross-sell : séquencement dépend des arbitrages priorité/capacité/owners. Neoscreen déploiement à partir d’octobre ; fin déploiement non fixée. Preuves de valeur selon disponibilité des cas et validation clients/métiers. Campagnes performance octobre–mi-décembre et février–avril. Les contenus run post-déploiement ne sont pas une nouvelle campagne.'

# A13 / detailed rituals roadmap
s=page('Annexe A13 / Roadmap détaillée','Rituels, série éditoriale et fils rouges.',subtitle='Août / vacances : rythme réduit. Les fils rouges ne sont pas des campagnes ponctuelles.',appendix=True)
timehead(s)
rit=[('All Staff France','staff'),('La Leytonienne','mag'),('Leadership au féminin','female'),('RDV / RSE mensuels','monthly'),('LinkedIn corporate','run'),('Temps forts collectifs','events')]
for i,(h,k) in enumerate(rit):
 y=4.66+i*.74;text(s,.95,y-.22,4.6,.87,h,22,N,True,leading=27)
 if k=='staff':
  for m in [3,6,9,11]:mark(s,y,m,N)
 elif k=='mag':mark(s,y,6,O);bar(s,y,10,12,N)
 elif k=='female':
  for m in [4,5,6,7,8]:mark(s,y,m,O)
 elif k=='monthly':
  for m in range(12):mark(s,y,m,N)
 elif k=='run':bar(s,y,0,12,N)
 else:
  bar(s,y+.15,0,12,G,True);text(s,TX+.16,y-.34,12.5,.45,'Christmas Party / séminaire : dates à confirmer',19,G)
 rule(s,.95,y+.34,18.1)
text(s,.95,9.17,18.1,.9,'Leytonienne : S1 fin janvier ; S2 préparation mai–juin → publication fin juillet, hors FY.\nLeadership : 4 capsules nov.–fév. + mash-up autour du 8 mars. All Staff fin exercice : date à confirmer.',22,G,leading=29)
s.notes_slide.notes_text_frame.text='Minimum 1 rendez-vous collaborateurs/mois et 1 activation RSE interne/mois. Ces repères ne supposent pas un post LinkedIn RSE/RH chaque mois. LinkedIn corporate fil rouge annuel ; seul calendrier octobre–juin fourni, pas de sujets inventés juillet–septembre. La préparation S2 en mai–juin n’est pas la publication, qui intervient fin juillet hors FY. Christmas Party et séminaire dates à confirmer ; on ne déduit pas un engagement événementiel de la seule position de Noël dans le calendrier LinkedIn.'

# A14-15 / governance. Four fields per large programme, no asserted individual owner.
def govpage(section,title,rows,foot):
 s=page(section,title,subtitle='Responsabilités fonctionnelles fournies ; owners opérationnels et validations à confirmer.',appendix=True)
 for i,(h,owner,contrib,comm,arb) in enumerate(rows):
  y=3.85+i*1.78;label(s,.95,y,h,O,17.9)
  text(s,.95,y+.58,5.45,1.18,'FOND : '+owner+'\nAVEC : '+contrib,21,N,leading=27)
  text(s,6.9,y+.58,6.2,1.18,'COMMUNICATION\n'+comm,21,G,leading=27)
  text(s,13.65,y+.58,5.4,1.18,'ARBITRAGE\n'+arb,21,G,leading=27)
  if i<2:rule(s,.95,y+1.65,18.1)
 takeaway(s,foot)
 return s
s=govpage('Annexe A14 / Gouvernance','Transformation : fond, orchestration et arbitrages.',[
 ('ADOPTION DE L’IA','Direction / DSI / Juridique','RH Formation / métiers','Parcours, ressources, relais','Owners par volet / Zoom IA'),
 ('TOOLBOX / ADOPTION DES OUTILS','DSI / métiers, selon outil','RH Formation selon besoin','Hub, fiches, standard, relais','Inventaire / validations / séquence'),
 ('BUSINESS PARTNER & CROSS-SELL','Marketing / métiers / Business','CS / experts concernés','Expérience, branding, récit, capitalisation','Périmètre BP / owner / sourcing')],
 'Neoscreen / intranet : owner projet et responsable technique à confirmer ; Communication : expérience et contenus.')
s.notes_slide.notes_text_frame.text='Le owner opérationnel unique n’est pas défini dans le brief. Les fonctions ici sont les propriétaires du fond à mobiliser, pas un RACI validé. Business Partner est un nom de travail, périmètre/owner à arbitrer. La répartition IA respecte strictement Direction sens, DSI outils/anonymisation, Juridique cadre, RH Formation compétences, métiers cas d’usage, Communication orchestration.'
s=govpage('Annexe A15 / Gouvernance & périmètres','Preuves, performance, People : travailler avec le fond.',[
 ('PREUVES & MONTÉE EN GAMME','Marketing / Business / CS','Métiers / clients / Groupe','Storytelling, formats, cohérence','Cas / droits / offre / calendrier Groupe'),
 ('PERFORMANCE','Direction / Finance','Marketing / Business / métiers','Déclinaisons validées / valorisation','Validation chiffres / critères'),
 ('CULTURE, ENGAGEMENT & IMAGE','RH / métiers','Managers / CSE / Sport Business','Orchestration / série / image corporate','Workelo / événements / validation')],
 'France Communication : corporate. Marketing : LinkedIn business. Groupe : projets Groupe. Sport Business : hospitalités sélectionnées.')
s.notes_slide.notes_text_frame.text='Périmètres : Communication France storytelling/branding/formats/cohérence/amplification/capitalisation ; Marketing, Business, Customer Success et métiers fond et stratégie commerciale ; Direction arbitrages ; Finance chiffres validés ; RH programmes RH ; DSI outils ; Juridique cadre ; Groupe projets/calendriers Groupe ; Sport Business hospitalités et grands temps forts sélectionnés, owner opérationnel non précisé. Neoscreen/intranet : responsable technique et owner projet à confirmer, Communication expérience et contenu, arbitrage après revue de production/prestataire.'

# A16-17 / actionable proposed measures
measures=[
 ('Adoption de l’IA','Contribuer à des usages compris et responsables.','Participation aux formations ;\nusages déclarés.','RH Formation ; retours métiers / enquête, à valider.'),
 ('Toolbox / adoption des outils','Faciliter l’accès et l’usage des outils utiles.','Consultations des ressources ; usages des outils.','Statistiques hub après lancement ; DSI / métiers.'),
 ('Business Partner & cross-sell','Équiper et faire capitaliser les collaborations.','Usage du kit ; cas qualifiés / capitalisés.','Données hub disponibles ; registre de sourcing.'),
 ('Preuves & montée en gamme','Renforcer la crédibilité par des preuves de valeur.','Preuves validées ; performance des contenus vidéo.','Registre éditorial ; analytics des plateformes.'),
 ('Performance','Rendre compréhensibles chiffres et offres.','Portée des relais ; retours sur leur compréhension.','Canaux disponibles ; retours équipes / managers.'),
 ('Culture, engagement & image','Contribuer à la mobilisation et à la visibilité du collectif.','Participation aux rendez-vous ; engagement corporate.','RH / organisateurs ; analytics réseaux corporate.'),
 ('Écosystème interne','Faciliter l’accès à une information utile.','Usage des canaux ; contribution / retour utilisateurs.','Statistiques disponibles après déploiement / lancement.')]
for n,rr in enumerate([measures[:3],measures[3:]],16):
 s=page(f'Annexe A{n} / Mesure','Effets recherchés, indicateurs et sources.',subtitle='Indicateurs proposés à valider selon disponibilité des données. Aucune cible numérique ajoutée.',appendix=True)
 for i,(h,e,ind,source) in enumerate(rr):
  y=3.85+i*(1.32 if len(rr)==4 else 1.65)
  label(s,.95,y,h,O,18)
  text(s,.95,y+.47,5.45,.84,e,22,N,True,leading=27)
  text(s,6.8,y+.47,5.65,.84,ind,22,G,leading=27)
  text(s,13.0,y+.47,6.03,.84,source,21,G,leading=26)
  if i<len(rr)-1:rule(s,.95,y+(1.22 if len(rr)==4 else 1.5),18.1)
 takeaway(s,'Repères d’activité : ≥ 4 vidéos cross-sell / FY ; ≥ 1 RDV et ≥ 1 activation RSE / mois ; 4 All Staff.')
 s.notes_slide.notes_text_frame.text='Un effet recherché et 1 à 2 indicateurs par grand ensemble. Sources proposées à mobiliser avec les owners, accès et granularité à confirmer : elles ne sont pas présentées comme existantes ou déjà exploitées. Les usages/participations ne suffisent pas à attribuer une évolution à Communication seule. Pas de cible quantitative ajoutée ; les volumes fournis sont une discipline de production, pas une preuve d’impact. Sources avant intranet : retours métiers, données canaux existants, RH et registre de production ; données intranet uniquement après lancement réel.'

# LinkedIn: verbatim transcription of the user image. No topics inferred.
TERR={'Corporate & innovation':('E6EBF0',N,8),'Leadership au féminin':('F3E8F7','AC2CCB',5),'RSE':('EDF5E5','4A9225',8),'Vie interne & RH':('FCECF1','D93662',8),'Sport':('E2F4F7','00869A',5)}
CAL=[
 ('OCT.','RSE','15/10 · Octobre Rose','Mobilisation, 1 km = 1 €, actions agences, webinaire','planifié'),
 ('OCT.','Corporate & innovation','22/10 · IA','Former aujourd’hui les métiers de demain','planifié'),
 ('OCT.','Sport','29/10 · HYROX','La tendance HYROX chez les collaborateurs Leyton','planifié'),
 ('NOV.','Corporate & innovation','Post All Staff','3 enseignements / convictions','planifié'),
 ('NOV.','Sport','~13/11 · Sport Business','France × Afrique du Sud','planifié'),
 ('NOV.','Vie interne & RH','Event RH','Movember','planifié'),
 ('NOV.','Leadership au féminin','26/11 · Leadership au féminin #1','','planifié'),
 ('DÉC.','RSE','RSE du mois','','planifié'),
 ('DÉC.','Leadership au féminin','Leadership au féminin #2','Le leadership sans nécessairement manager','planifié'),
 ('DÉC.','Vie interne & RH','Event RH','Soirée de Noël','planifié'),
 ('DÉC.','Corporate & innovation','Vœux Leyton','','planifié'),
 ('JAN.','Corporate & innovation','Rétrospective','Ce que nous emportons de 2026 en 2027','planifié'),
 ('JAN.','RSE','RSE','Collecte de vêtements','planifié'),
 ('JAN.','Vie interne & RH','Event RH du mois','','planifié'),
 ('JAN.','Leadership au féminin','Leadership au féminin #3','Les compétences de demain : pas seulement techniques','planifié'),
 ('FÉV.','Corporate & innovation','Challenge DigiCraft','L’innovation en train de se faire','planifié'),
 ('FÉV.','RSE','RSE du mois','','planifié'),
 ('FÉV.','Vie interne & RH','Event RH du mois','','planifié'),
 ('FÉV.','Leadership au féminin','Leadership au féminin #4','Mobilité & évolution professionnelle','planifié'),
 ('MARS','Leadership au féminin','Semaine du 8 mars','Mash-up Leadership au féminin','planifié'),
 ('MARS','Sport','Semi-marathon','Event sportif','À confirmer'),
 ('MARS','Vie interne & RH','Event RH du mois','','planifié'),
 ('MARS','RSE','RSE du mois','','planifié'),
 ('AVR.','Corporate & innovation','Post All Staff','Enseignement / conviction','planifié'),
 ('AVR.','RSE','RSE du mois','','planifié'),
 ('AVR.','Vie interne & RH','Event RH du mois','','planifié'),
 ('AVR.','Corporate & innovation','Leyton 2030','Sujet corporate','À définir'),
 ('MAI','Sport','Sport Business','Hospitalité / temps fort PSG, OL ou Stade Toulousain','À sélectionner'),
 ('MAI','RSE','RSE du mois','','planifié'),
 ('MAI','Vie interne & RH','Event RH du mois','','planifié'),
 ('MAI','Corporate & innovation','Sujet corporate','','À définir'),
 ('JUIN','RSE','RSE du mois','','planifié'),
 ('JUIN','Vie interne & RH','Event RH du mois','','planifié'),
 ('JUIN','Sport','Triathlon','Event sportif','À confirmer')]
OPP=[('FÉV.','Sport','21/02 · Sport Business','France × Écosse','Opportunité'),('MARS','Sport','Actualité sport','Uniquement si réellement stratégique','Opportunité')]
assert len(CAL)==34
for name,(_,_,count) in TERR.items():assert sum(r[1]==name for r in CAL)==count
assert sum(r[4]!='planifié' for r in CAL)==5
(ROOT/'calendrier_linkedin.json').write_text(json.dumps({'source':'Image fournie par utilisateur : Calendrier éditorial LinkedIn, Exercice 29, octobre à juin. Transcription manuelle sans ajout de sujets.','entries':CAL,'opportunities':OPP},ensure_ascii=False,indent=2))
s=page('Annexe A18 / LinkedIn corporate','Un calendrier éditorial, cinq territoires de marque.',subtitle='Source : calendrier joint • octobre → juin • 34 entrées, dont 5 à confirmer / définir / sélectionner.',appendix=True)
for i,(h,(fill,color,count)) in enumerate(TERR.items()):
 y=3.85+i*.89;box(s,.95,y,8.6,.69,fill);text(s,1.2,y+.13,7.9,.42,f'{h} · {count}',24,color,True)
text(s,10.25,3.95,8.7,2.14,'Cadence du calendrier :\n3 à 4 entrées par mois.\n29 sans réserve explicite ; 5 conditionnelles.',27,N,True,leading=36)
text(s,10.25,6.22,8.7,1.47,'Deux opportunités sport en plus,\nhors des 34 entrées : février et mars.\nMaximum un post par événement.',25,G,leading=34)
text(s,10.25,8.04,8.7,.93,'Sport collaborateurs / CSE ≠ Sport Business.\nAucune couverture systématique.',24,N,True,leading=32)
takeaway(s,'Ce calendrier ne représente ni le LinkedIn business Marketing, ni l’ensemble du run interne RH / RSE.')
s.notes_slide.notes_text_frame.text='Les 34 entrées reprennent exactement les 5 comptes de la légende de l’image : Corporate & innovation 8 ; Leadership 5 ; RSE 8 ; Vie interne & RH 8 ; Sport 5. Les cartes en pointillés sont incluses dans ce total : semi-marathon, Leyton 2030, Sport Business mai, sujet corporate mai, triathlon. Les deux opportunités orange sont séparées. La première capsule ne reçoit pas d’angle inventé. Aucun sujet juillet–septembre ajouté. Le post IA du 22 octobre ne fixe pas la date du Zoom IA.'
for n,months in enumerate([['OCT.','NOV.','DÉC.'],['JAN.','FÉV.','MARS'],['AVR.','MAI','JUIN']],19):
 s=page(f'Annexe A{n} / Calendrier LinkedIn',' / '.join(months)+' — sujets du calendrier fourni.',subtitle='Contenus et statuts transcrits de l’image ; aucune nouvelle proposition éditoriale.',appendix=True)
 for j,m in enumerate(months):
  x=.95+j*6.15;text(s,x,3.68,5.8,.58,m,30,N,True)
  for k,r in enumerate([r for r in CAL if r[0]==m]):
   _,territory,head,body,status=r;fill,color,count=TERR[territory];y=4.4+k*1.38
   sh=box(s,x,y,5.78,1.30,fill if status=='planifié' else W,line=color if status!='planifié' else None)
   if status!='planifié':sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
   text(s,x+.16,y+.1,5.46,.49,head,21,color,True,leading=24)
   content=body+((' · '+status) if body else status) if status!='planifié' else body
   if content:text(s,x+.16,y+.61,5.46,.70,content,18,G,leading=21)
 if n==21:
  text(s,.95,10.02,18.1,.40,'OPPORTUNITÉS : 21/02 Sport Business — France × Écosse ; mars — actualité sport uniquement si réellement stratégique.',18,O,True)
 else:
  text(s,.95,10.02,18.1,.40,'Cadence et thèmes du calendrier source ; dates des événements à confirmer lorsqu’indiqué.',19,G)
 s.notes_slide.notes_text_frame.text='Transcription fidèle du calendrier LinkedIn fourni. '+json.dumps([r for r in CAL if r[0] in months],ensure_ascii=False)+'\nOpportunités hors 34 entrées : 21/02 Sport Business France × Écosse ; mars actualité sport uniquement si réellement stratégique. Aucun sujet ajouté. Les intitulés et détails complets figurent aussi dans calendrier_linkedin.json.'

# Final ordering: 13 newly refactored main slides; 11 retained V3 annexes; 10 new annexes.
ORDER=MAIN+ANNEX_IDS+EXTRA
lst=P.slides._sldIdLst
for sid in list(lst):lst.remove(sid)
for sid in ORDER:lst.append(sid)
# Remove unused relationships (not reused original slides or generated ones).
used={sid.rId for sid in ORDER}
for sid in ORIGINAL_IDS:
 if sid.rId not in used:P.part.drop_rel(sid.rId)
assert len(P.slides)==34
# Renumber existing footers without altering original V3 file.
for i,s in enumerate(P.slides,1):
 for sh in s.shapes:
  if sh.has_text_frame and sh.left>Inches(17.5) and sh.top>Inches(10.5) and sh.text.strip().isdigit():para_replace(s,sh.text,f'{i:02d}')
 P.core_properties.title='Plan de Communication France FY2026/27 — V4'
 P.core_properties.subject='Deck de direction : choix, temporalité, gouvernance, mesure et arbitrages'
 P.core_properties.author='Communication France — Leyton'
OUT.mkdir(parents=True,exist_ok=True)
P.save(OUT/'Plan_Communication_France_FY2026-27_Leyton_V4.pptx')
manifest=[];geo=[]
for i,s in enumerate(P.slides,1):
 blocks=[]
 for sh in s.shapes:
  if sh.has_text_frame and sh.text.strip():
   blocks.append(sh.text)
   if sh.text_frame.paragraphs and sh.text_frame.paragraphs[0].runs:
    size=sh.text_frame.paragraphs[0].runs[0].font.size or sh.text_frame.paragraphs[0].font.size
    if size:geo.append({'slide':i,'text':sh.text,'x':sh.left/12700,'y':sh.top/12700,'w':sh.width/12700,'h':sh.height/12700,'size':size.pt})
 manifest.append({'slide':i,'section':'main' if i<=13 else 'appendix','text':blocks,'notes':s.notes_slide.notes_text_frame.text})
(OUT/'content_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
(OUT/'text_geometry.json').write_text(json.dumps(geo,ensure_ascii=False,indent=2))
print('Created V4:',len(P.slides),'slides (13 main + 21 annexes). Base V3 SHA256 unchanged:',BASE_HASH)
